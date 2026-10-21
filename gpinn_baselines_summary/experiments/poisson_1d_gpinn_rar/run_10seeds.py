"""Run and audit both Poisson gPINN+RAR cases over multiple random seeds."""

from __future__ import annotations

import argparse
import csv
import json
import os
import queue
import subprocess
import sys
import threading
from pathlib import Path

import numpy as np
import torch

from .cases import CASES
from .run import METHOD_NAMES, PROTOCOLS


CASE_LABELS = {
    "boundary_layer": "1D Poisson (boundary-layer)",
    "steep": "1D Poisson (steep)",
}


def parse_seeds(value: str) -> tuple[int, ...]:
    result: list[int] = []
    for item in value.split(","):
        token = item.strip()
        if not token:
            continue
        if "-" in token:
            start, stop = (int(part) for part in token.split("-", 1))
            if stop < start:
                raise ValueError(f"descending seed range: {token}")
            result.extend(range(start, stop + 1))
        else:
            result.append(int(token))
    if not result or len(result) != len(set(result)):
        raise ValueError("seeds must be non-empty and unique")
    return tuple(result)


def parse_gpus(value: str) -> tuple[str, ...]:
    if value.lower() == "cpu":
        return ()
    if value.lower() == "auto":
        visible = os.environ.get("CUDA_VISIBLE_DEVICES", "").strip()
        if visible and visible != "-1":
            return tuple(item.strip() for item in visible.split(",") if item.strip())
        if torch.cuda.is_available():
            return tuple(str(index) for index in range(torch.cuda.device_count()))
        return ()
    result = tuple(item.strip() for item in value.split(",") if item.strip())
    if not result or any(not item.isdigit() for item in result):
        raise ValueError("--gpus expects auto, cpu, or comma-separated GPU IDs")
    return result


def result_dir(root: Path, mode: str, case: str, seed: int) -> Path:
    return root / mode / case / f"seed_{seed}"


def is_complete(root: Path, mode: str, case: str, seed: int, protocol: str) -> bool:
    path = result_dir(root, mode, case, seed) / "metrics.json"
    if not path.exists():
        return False
    try:
        metrics = json.loads(path.read_text(encoding="utf-8"))
        return (
            metrics.get("method") == METHOD_NAMES[protocol]
            and metrics.get("protocol") == protocol
            and metrics.get("mode") == mode
            and metrics.get("case") == case
            and int(metrics.get("seed", -1)) == seed
            and (mode != "full" or bool(metrics.get("full_training_completed")))
        )
    except (OSError, ValueError, json.JSONDecodeError):
        return False


def run_task(
    args: argparse.Namespace,
    extra: list[str],
    case: str,
    seed: int,
    gpu: str | None,
    threads: int,
) -> tuple[str, int, int]:
    directory = result_dir(args.output_root, args.mode, case, seed)
    directory.mkdir(parents=True, exist_ok=True)
    command = [
        sys.executable,
        "-m",
        "experiments.poisson_1d_gpinn_rar.run",
        "--case",
        case,
        "--seed",
        str(seed),
        "--mode",
        args.mode,
        "--device",
        "cuda" if gpu is not None else "cpu",
        "--output-root",
        str(args.output_root),
        "--protocol",
        args.protocol,
        *extra,
    ]
    if not args.force and (directory / "last.pt").exists():
        command.append("--resume")
    environment = os.environ.copy()
    if gpu is not None:
        environment["CUDA_VISIBLE_DEVICES"] = gpu
    else:
        environment["CUDA_VISIBLE_DEVICES"] = ""
    for name in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS"):
        environment[name] = str(threads)
    environment["MPLCONFIGDIR"] = str(directory / ".mplconfig")
    with (directory / "launcher.log").open("a", encoding="utf-8") as log:
        log.write("\nCOMMAND: " + " ".join(command) + "\n")
        completed = subprocess.run(
            command,
            cwd=str(Path(__file__).resolve().parent.parent),
            env=environment,
            stdout=log,
            stderr=subprocess.STDOUT,
            check=False,
        )
    return case, seed, completed.returncode


def _audit_run(
    directory: Path,
    mode: str,
    case: str,
    seed: int,
    protocol: str,
) -> dict:
    required = (
        "config.json",
        "last.pt",
        "metrics.json",
        "prediction.npz",
        "reference.npz",
        "rar_points.npz",
        "rar_history.jsonl",
        "training_log.csv",
        "solution_error_rar.png",
        "absolute_error.png",
        "loss_curve.png",
    )
    missing = [name for name in required if not (directory / name).exists()]
    if missing:
        raise FileNotFoundError(f"{case}/seed_{seed}: missing {missing}")
    metrics = json.loads((directory / "metrics.json").read_text(encoding="utf-8"))
    config = json.loads((directory / "config.json").read_text(encoding="utf-8"))
    if config["hidden_layers"] != 6 or config["width"] != 128:
        raise AssertionError(f"{case}/seed_{seed}: network is not 6x128")
    expected_score = (
        "absolute_raw_pde_residual"
        if protocol == "burgers_rar_transfer"
        else "not_applicable"
    )
    if (
        metrics["method"] != METHOD_NAMES[protocol]
        or metrics["protocol"] != protocol
        or metrics["rar_score"] != expected_score
    ):
        raise AssertionError(f"{case}/seed_{seed}: method identity mismatch")
    if mode == "full" and not metrics["full_training_completed"]:
        raise AssertionError(f"{case}/seed_{seed}: incomplete full run")
    expected_points = int(config["initial_points"]) + int(config["rar_cycles"]) * int(
        config["rar_points_per_cycle"]
    )
    if int(metrics["collocation_points"]) != expected_points:
        raise AssertionError(f"{case}/seed_{seed}: RAR point count mismatch")
    history = [
        json.loads(line)
        for line in (directory / "rar_history.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    expected_cycles = list(range(1, int(config["rar_cycles"]) + 1))
    if [int(row["cycle"]) for row in history] != expected_cycles:
        raise AssertionError(f"{case}/seed_{seed}: RAR history is incomplete or duplicated")
    with np.load(directory / "prediction.npz") as data:
        prediction = data["prediction"]
    with np.load(directory / "reference.npz") as data:
        reference = data["reference"]
    if not np.isfinite(prediction).all() or not np.isfinite(reference).all():
        raise FloatingPointError(f"{case}/seed_{seed}: NaN or Inf")
    error = prediction - reference
    relative_l2 = float(np.linalg.norm(error) / np.linalg.norm(reference))
    max_abs = float(np.max(np.abs(error)))
    if not np.isclose(relative_l2, metrics["relative_l2"], rtol=1e-12, atol=1e-14):
        raise AssertionError(f"{case}/seed_{seed}: relative L2 cannot be reproduced")
    if not np.isclose(max_abs, metrics["max_abs"], rtol=1e-12, atol=1e-14):
        raise AssertionError(f"{case}/seed_{seed}: max error cannot be reproduced")
    return metrics


def write_summary(
    root: Path,
    mode: str,
    cases: tuple[str, ...],
    seeds: tuple[int, ...],
    protocol: str,
) -> None:
    rows = [
        _audit_run(result_dir(root, mode, case, seed), mode, case, seed, protocol)
        for case in cases
        for seed in seeds
    ]
    mode_root = root / mode
    fields = (
        "case",
        "seed",
        "method",
        "protocol",
        "architecture",
        "collocation_points",
        "rar_cycles_completed",
        "adam_steps_completed",
        "lbfgs_closure_evaluations",
        "relative_l2",
        "max_abs",
        "boundary_max_abs",
        "reference_type",
        "full_training_completed",
        "device",
    )
    with (mode_root / "results_summary.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    aggregate = []
    for case in cases:
        selected = [row for row in rows if row["case"] == case]
        l2 = np.asarray([row["relative_l2"] for row in selected], dtype=np.float64)
        maximum = np.asarray([row["max_abs"] for row in selected], dtype=np.float64)
        l2_std = float(np.std(l2, ddof=1)) if len(l2) > 1 else 0.0
        max_std = float(np.std(maximum, ddof=1)) if len(maximum) > 1 else 0.0
        aggregate.append(
            {
                "case": case,
                "case_name": CASE_LABELS[case],
                "protocol": protocol,
                "method": METHOD_NAMES[protocol],
                "seeds": ",".join(str(seed) for seed in seeds),
                "completed_runs": len(selected),
                "architecture": selected[0]["architecture"],
                "initial_points": json.loads(
                    (result_dir(root, mode, case, seeds[0]) / "config.json").read_text()
                )["initial_points"],
                "final_points": selected[0]["collocation_points"],
                "rar_cycles": selected[0]["rar_cycles_completed"],
                "relative_l2_mean": float(np.mean(l2)),
                "relative_l2_std": l2_std,
                "relative_l2_mean_std": f"{np.mean(l2):.6e} ± {l2_std:.6e}",
                "max_abs_mean": float(np.mean(maximum)),
                "max_abs_std": max_std,
                "max_abs_mean_std": f"{np.mean(maximum):.6e} ± {max_std:.6e}",
                "reference_type": "analytic",
                "all_full_training_completed": all(
                    bool(row["full_training_completed"]) for row in selected
                ),
            }
        )
    aggregate_fields = tuple(aggregate[0])
    with (mode_root / "results_mean_std.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=aggregate_fields)
        writer.writeheader()
        writer.writerows(aggregate)
    with (mode_root / "results_mean_std.md").open("w", encoding="utf-8") as handle:
        handle.write(f"# 1D Poisson {METHOD_NAMES[protocol]} 十种子统计\n\n")
        handle.write("标准差为样本标准差（ddof=1），只统计审计通过的实际结果。\n\n")
        handle.write("| 算例 | 完成数 | 相对 L2（mean ± std） | 最大误差（mean ± std） |\n")
        handle.write("|---|---:|---:|---:|\n")
        for row in aggregate:
            handle.write(
                f"| {row['case_name']} | {row['completed_runs']} | "
                f"{row['relative_l2_mean_std']} | {row['max_abs_mean_std']} |\n"
            )

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as error:
        raise RuntimeError("install openpyxl to write the required XLSX") from error
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "十种子汇总"
    headers = (
        "算例",
        "协议",
        "随机种子",
        "完成次数",
        "网络",
        "初始点数",
        "RAR后点数",
        "RAR轮数",
        "相对L2（mean ± std）",
        "最大误差（mean ± std）",
        "参考解",
        "完整训练完成",
    )
    sheet.append(headers)
    for row in aggregate:
        sheet.append(
            (
                row["case_name"],
                row["protocol"],
                row["seeds"],
                row["completed_runs"],
                row["architecture"],
                row["initial_points"],
                row["final_points"],
                row["rar_cycles"],
                row["relative_l2_mean_std"],
                row["max_abs_mean_std"],
                row["reference_type"],
                "是" if row["all_full_training_completed"] else "否",
            )
        )
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(
            45, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    sheet.freeze_panes = "A2"
    workbook.save(mode_root / "results_mean_std.xlsx")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seeds", default="0-9")
    parser.add_argument("--cases", default="boundary_layer,steep")
    parser.add_argument("--mode", choices=("smoke", "full"), required=True)
    parser.add_argument(
        "--protocol", choices=PROTOCOLS, default="burgers_rar_transfer"
    )
    parser.add_argument("--gpus", default="0,1,2,3,4,5,6,7")
    parser.add_argument("--cpus", type=int, default=32)
    parser.add_argument("--cpu-workers", type=int, default=1)
    parser.add_argument(
        "--output-root",
        type=Path,
        default=None,
    )
    parser.add_argument("--force", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> None:
    args, extra = build_parser().parse_known_args(argv)
    if args.output_root is None:
        args.output_root = (
            Path(__file__).resolve().parents[2]
            / "results"
            / f"poisson_1d_{args.protocol}_10seeds"
        )
    args.output_root = args.output_root.resolve()
    args.output_root.mkdir(parents=True, exist_ok=True)
    seeds = parse_seeds(args.seeds)
    cases = tuple(item.strip() for item in args.cases.split(",") if item.strip())
    if not cases or any(case not in CASES for case in cases):
        raise ValueError(f"cases must come from {sorted(CASES)}")
    gpus = parse_gpus(args.gpus)
    workers = tuple(gpus) if gpus else tuple(None for _ in range(args.cpu_workers))
    if not workers or args.cpus <= 0:
        raise ValueError("at least one worker and a positive CPU count are required")
    threads = max(1, args.cpus // len(workers))
    pending = [
        (case, seed)
        for case in cases
        for seed in seeds
        if args.force
        or not is_complete(args.output_root, args.mode, case, seed, args.protocol)
    ]
    print(
        json.dumps(
            {
                "method": METHOD_NAMES[args.protocol],
                "protocol": args.protocol,
                "mode": args.mode,
                "cases": cases,
                "seeds": seeds,
                "gpus": gpus,
                "pending": [f"{case}:{seed}" for case, seed in pending],
                "threads_per_worker": threads,
                "output_root": str(args.output_root),
            },
            indent=2,
        ),
        flush=True,
    )
    tasks: queue.Queue[tuple[str, int]] = queue.Queue()
    for task in pending:
        tasks.put(task)
    failures: list[tuple[str, int, int]] = []
    lock = threading.Lock()

    def worker(gpu: str | None) -> None:
        label = f"GPU {gpu}" if gpu is not None else "CPU"
        while True:
            try:
                case, seed = tasks.get_nowait()
            except queue.Empty:
                return
            with lock:
                print(f"[{label}] {case} seed={seed} started", flush=True)
            outcome = run_task(args, extra, case, seed, gpu, threads)
            if outcome[2] != 0:
                failures.append(outcome)
            with lock:
                status = "completed" if outcome[2] == 0 else f"failed({outcome[2]})"
                print(f"[{label}] {case} seed={seed} {status}", flush=True)
            tasks.task_done()

    threads_list = [threading.Thread(target=worker, args=(gpu,)) for gpu in workers]
    for thread in threads_list:
        thread.start()
    for thread in threads_list:
        thread.join()
    if failures:
        raise RuntimeError(f"failed tasks: {failures}")
    write_summary(args.output_root, args.mode, cases, seeds, args.protocol)
    print(f"[done] {args.output_root / args.mode / 'results_mean_std.xlsx'}", flush=True)


if __name__ == "__main__":
    main()
