from __future__ import annotations

import numpy as np
import torch
import json

from gpinn_baselines.cases import REGISTRY, create_case
from gpinn_baselines.aggregate import METRIC_FIELDS, _write_xlsx, aggregate_rows
from gpinn_baselines.config import CASE_NAMES, build_config
from gpinn_baselines.metrics import cavity_metrics, scalar_metrics
from gpinn_baselines.network import MLP
from gpinn_baselines.references import interpolate_reference
from gpinn_baselines.seeds import parse_seed_spec
from gpinn_baselines.trainer import _cpu_byte_rng_states


def test_all_thirteen_cases_are_registered():
    assert len(CASE_NAMES) == 13
    assert tuple(REGISTRY) == CASE_NAMES


def test_ten_seed_spec_and_mean_std_aggregation(tmp_path):
    seeds = parse_seed_spec("0-9")
    assert seeds == tuple(range(10))
    rows = []
    for case in CASE_NAMES:
        for seed in seeds:
            row = {
                "case": case,
                "seed": seed,
                "network": "6x128 tanh (float64)",
                "n_residual": 100,
                "n_boundary": 20,
                "n_initial": 20,
                "adam_iterations": 20000,
                "lbfgs_iterations": 1000,
                "relative_l2": float(seed),
                "max_abs": float(2 * seed),
                "reference_type": "analytic",
                "full_training_completed": True,
            }
            row.update({metric: None for metric in METRIC_FIELDS[2:]})
            if case == "lid_driven_cavity":
                row.update({metric: float(seed + 1) for metric in METRIC_FIELDS[2:]})
            rows.append(row)
    summary = aggregate_rows(rows, expected_seeds=seeds, require_complete=True)
    assert len(summary) == 13
    assert summary[0]["completed_runs"] == 10
    assert np.isclose(summary[0]["relative_l2_mean"], 4.5)
    assert np.isclose(
        summary[0]["relative_l2_std"], np.std(np.arange(10), ddof=1)
    )
    assert summary[0]["relative_l2_mean_std"].count("±") == 1
    poisson_only = aggregate_rows(
        rows,
        expected_seeds=seeds,
        require_complete=True,
        case_names=CASE_NAMES[:2],
    )
    assert [row["case"] for row in poisson_only] == list(CASE_NAMES[:2])
    workbook_path = tmp_path / "results_mean_std.xlsx"
    _write_xlsx(summary, workbook_path)
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True)
    sheet = workbook["十种子汇总"]
    assert sheet.max_row == 14
    assert sheet["F2"].value == summary[0]["relative_l2_mean_std"]
    workbook.close()


def test_network_contract():
    model = MLP(2, 1, [-1.0, 0.0], [1.0, 1.0])
    assert len(model.hidden) == 6
    assert all(layer.out_features == 128 for layer in model.hidden)
    assert model(torch.zeros((3, 2), dtype=torch.float64)).shape == (3, 1)


def test_poisson_1d_original_gpinn_protocol_and_hard_constraints():
    for name in ("poisson_1d_boundary_layer", "poisson_1d_steep"):
        config = build_config(name, "full", 0)
        assert config["sampling"] == {
            "n_residual": 500,
            "n_boundary": 0,
            "n_initial": 0,
            "distribution": "latin_hypercube",
        }
        assert config["optimizer"]["learning_rate"] == 1.0e-3
        assert config["optimizer"]["lr_decay_rate"] == 1.0
        assert config["optimizer"]["adam_iterations"] == 20000
        assert config["optimizer"]["lbfgs_iterations"] == 0
        assert config["loss_weights"]["gradient"] == 0.01

        case = create_case(config)
        batch = case.sample_training(np.random.default_rng(0))
        assert tuple(batch) == ("residual",)
        model = MLP(1, 1, config["lower"], config["upper"])
        model.set_output_transform(case.output_transform)
        endpoints = torch.tensor([[-1.0], [1.0]], dtype=torch.float64)
        np.testing.assert_allclose(
            model(endpoints).detach().numpy(),
            case.exact_numpy(endpoints.numpy()),
            rtol=0.0,
            atol=1.0e-13,
        )


def test_burgers_original_repository_protocol_and_hard_constraints():
    config = build_config("burgers_1d", "full", 0)
    assert config["sampling"] == {
        "n_residual": 1500,
        "n_boundary": 0,
        "n_initial": 0,
        "distribution": "latin_hypercube",
    }
    assert config["optimizer"]["learning_rate"] == 1.0e-3
    assert config["optimizer"]["lr_decay_rate"] == 1.0
    assert config["optimizer"]["adam_iterations"] == 20000
    assert config["optimizer"]["lbfgs_iterations"] == 15000
    assert config["loss_weights"]["gradient"] == 1.0e-4

    case = create_case(config)
    batch = case.sample_training(np.random.default_rng(0))
    assert tuple(batch) == ("residual",)
    model = MLP(2, 1, config["lower"], config["upper"])
    model.set_output_transform(case.output_transform)
    x = torch.linspace(-1.0, 1.0, 17, dtype=torch.float64)[:, None]
    initial = torch.column_stack((x[:, 0], torch.zeros(len(x), dtype=torch.float64)))
    np.testing.assert_allclose(
        model(initial).detach().numpy()[:, 0],
        -np.sin(np.pi * x.numpy()[:, 0]),
        atol=1.0e-13,
    )
    times = torch.linspace(0.0, 1.0, 11, dtype=torch.float64)
    left = torch.column_stack((-torch.ones_like(times), times))
    right = torch.column_stack((torch.ones_like(times), times))
    np.testing.assert_allclose(model(left).detach().numpy(), 0.0, atol=1.0e-13)
    np.testing.assert_allclose(model(right).detach().numpy(), 0.0, atol=1.0e-13)


def test_allen_cahn_contract_and_periodic_pairs():
    config = build_config("allen_cahn_1d", "smoke", 0)
    assert config["parameters"]["diffusion"] == 1.0e-4
    case = create_case(config)
    batch = case.sample_training(np.random.default_rng(0))
    np.testing.assert_array_equal(batch["periodic_left"][:, 1], batch["periodic_right"][:, 1])
    assert np.all(batch["periodic_left"][:, 0] == -1.0)
    assert np.all(batch["periodic_right"][:, 0] == 1.0)


def test_metric_definitions():
    reference = np.asarray([[1.0], [2.0]])
    prediction = np.asarray([[1.5], [1.0]])
    values = scalar_metrics(prediction, reference)
    assert np.isclose(values["relative_l2"], np.sqrt(1.25) / np.sqrt(5.0))
    assert values["max_abs"] == 1.0

    velocity_reference = np.asarray([[1.0, 0.0], [0.0, 1.0]])
    velocity_prediction = np.asarray([[1.0, 0.5, 9.0], [0.0, 0.0, -2.0]])
    cavity = cavity_metrics(velocity_prediction, velocity_reference)
    assert np.isclose(cavity["velocity_relative_l2"], np.sqrt(1.25) / np.sqrt(2.0))
    assert cavity["velocity_max_pointwise_l2"] == 1.0


def test_allen_cahn_periodic_reference_accepts_x_equals_one(tmp_path):
    x = np.linspace(-1.0, 1.0, 8, endpoint=False)
    t = np.asarray([0.0, 1.0])
    u = np.vstack((np.cos(np.pi * x), 2.0 * np.cos(np.pi * x)))
    path = tmp_path / "reference.npz"
    np.savez_compressed(
        path,
        x=x,
        t=t,
        u=u,
        metadata=np.asarray(json.dumps({"case": "allen_cahn_1d"})),
    )
    points = np.asarray([[-1.0, 0.0], [1.0, 0.0], [1.0, 1.0]])
    values, _ = interpolate_reference(path, "allen_cahn_1d", points)
    np.testing.assert_allclose(values[:, 0], [-1.0, -1.0, -2.0])


def test_checkpoint_cuda_rng_states_are_normalized_for_restore():
    states = _cpu_byte_rng_states([torch.tensor([1, 2, 3], dtype=torch.int64)])
    assert len(states) == 1
    assert states[0].device.type == "cpu"
    assert states[0].dtype == torch.uint8
    assert states[0].is_contiguous()
